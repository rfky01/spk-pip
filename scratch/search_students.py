import openpyxl

file_path = r"D:\KULIAH\Joki\Annisa\skripsi\BAB IV\daftar_pd-SMP TUNAS BANGSA-2026-10-17 05_49_36.xlsx"
wb = openpyxl.load_workbook(file_path)
ws = wb.active

students = []
for row in range(7, ws.max_row + 1):
    no = ws.cell(row, 1).value
    nama = ws.cell(row, 2).value
    nipd = ws.cell(row, 3).value
    jk = ws.cell(row, 4).value
    nisn = ws.cell(row, 5).value
    alamat = ws.cell(row, 10).value
    hp = ws.cell(row, 20).value
    telp = ws.cell(row, 19).value
    gaji_ayah = ws.cell(row, 29).value
    gaji_ibu = ws.cell(row, 35).value
    sekolah_asal = ws.cell(row, 57).value
    jarak = ws.cell(row, 66).value
    
    if nama:
        students.append({
            'row': row,
            'no': no,
            'nama': str(nama).strip(),
            'jk': jk,
            'nisn': str(nisn).strip() if nisn else '',
            'alamat': alamat,
            'hp': hp,
            'telp': telp,
            'gaji_ayah': gaji_ayah,
            'gaji_ibu': gaji_ibu,
            'sekolah_asal': sekolah_asal,
            'jarak': jarak
        })

print(f"Total students loaded from Excel: {len(students)}")

target_names = [
    "rahmat padil peratama",
    "abizar Nur ramadhan",
    "yusup Baihaqi",
    "ketut evi yani",
    "mita sari",
    "bima Syarif al bukhory",
    "abi Ibnu Mas'ud",
    "gede radit saputra",
    "reza agus pertama",
    "nyoman mahesa purna"
]

import difflib

for i, t in enumerate(target_names, 1):
    found = None
    t_clean = t.lower().replace("'", "").replace(" ", "")
    for s in students:
        s_clean = s['nama'].lower().replace("'", "").replace(" ", "")
        if t_clean in s_clean or s_clean in t_clean or difflib.SequenceMatcher(None, t_clean, s_clean).ratio() > 0.7:
            found = s
            break
    if found:
        print(f"{i}. Found target '{t}' -> '{found['nama']}' (NISN: {found['nisn']}, Asal: {found['sekolah_asal']}, HP: {found['hp']}, GajiAyah: {found['gaji_ayah']}, Jarak: {found['jarak']})")
    else:
        print(f"{i}. NOT FOUND: '{t}'")
