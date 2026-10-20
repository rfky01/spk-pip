import docx

doc = docx.Document(r'D:\KULIAH\Joki\Annisa\(ok) reNEWok Proposal BAB I-III Annisa belum fix 2 (1) (1) - UPDATED.docx')
print('=== VERIFICATION OF UPDATED DOCUMENT ===')
print('Total paragraphs:', len(doc.paragraphs))
print('Total tables:', len(doc.tables))

# 1. Check Table 22
t22 = doc.tables[21]
print('Table 22 rows:', len(t22.rows))
print('Table 22 Row 1:', [c.text.strip() for c in t22.rows[1].cells])
print('Table 22 Row 20:', [c.text.strip() for c in t22.rows[20].cells])
print('Table 22 Row 30:', [c.text.strip() for c in t22.rows[30].cells])

# 2. Check for unwanted strings
issues = []
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if '24 siswa' in t: issues.append(f'P{i} has "24 siswa": {t[:80]}')
    if '27 siswa' in t: issues.append(f'P{i} has "27 siswa": {t[:80]}')
    if '10 siswa' in t: issues.append(f'P{i} has "10 siswa": {t[:80]}')
    if 'Tabel 4.20' in t: issues.append(f'P{i} has "Tabel 4.20": {t[:80]}')
    if 'Hampilan' in t: issues.append(f'P{i} has "Hampilan": {t[:80]}')

if not issues:
    print('ALL INCONSISTENT STRINGS FIXED! No issues found.')
else:
    for iss in issues:
        print('ISSUE:', iss)

# 3. Check Figure captions
print('\n=== FIGURE CAPTIONS ===')
for i, p in enumerate(doc.paragraphs):
    if 'Gambar 4.' in p.text:
        print(f'P{i}: {p.text.strip()}')

# 4. Check Table titles
print('\n=== TABLE TITLES IN BAB IV ===')
for i, p in enumerate(doc.paragraphs):
    if 'Tabel 4.' in p.text:
        print(f'P{i}: {p.text.strip()}')

