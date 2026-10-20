import openpyxl

wb = openpyxl.load_workbook(r"D:\KULIAH\Joki\Annisa\skripsi\BAB IV\daftar_pd-SMP TUNAS BANGSA-2026-10-17 05_49_36.xlsx")
ws = wb.active

target_rows = {
    "1. RAHMAD VADIL PRATAMA": 93,
    "2. ABIZHAR NURRAKHMADAN": 8,
    "3. YUSUF BAIHAQI": 115,
    "4. Ketut Evi Yani": 60,
    "5. Mita Sari": 75,
    "6. BIMA SYARIF AL BUKHORI": 25,
    "7. ABI IBNU MAS'UD": 7,
    "8. Gede Radit Saputra": 46,
    "9. REZA AGUS PRATAMA": 98,
    "10. Nyoman Mahesa Purna": 88
}

for label, r in target_rows.items():
    nama = ws.cell(r, 2).value
    nisn = ws.cell(r, 5).value
    jk = ws.cell(r, 4).value
    hp = ws.cell(r, 20).value
    telp = ws.cell(r, 19).value
    asal = ws.cell(r, 57).value
    alamat = ws.cell(r, 10).value
    rt = ws.cell(r, 11).value
    rw = ws.cell(r, 12).value
    dusun = ws.cell(r, 13).value
    kel = ws.cell(r, 14).value
    kec = ws.cell(r, 15).value
    gaji_ayah = ws.cell(r, 29).value
    gaji_ibu = ws.cell(r, 35).value
    tinggal = ws.cell(r, 17).value
    jarak = ws.cell(r, 66).value
    
    print(f"\n{label}:")
    print(f"  NISN: {nisn}, JK: {jk}, Asal: {asal}, HP: {hp}, Telp: {telp}")
    print(f"  Alamat: {alamat}, RT/RW: {rt}/{rw}, Dusun: {dusun}, Kel: {kel}, Kec: {kec}")
    print(f"  GajiAyah: {gaji_ayah}, GajiIbu: {gaji_ibu}, Tinggal: {tinggal}, Jarak: {jarak}")
