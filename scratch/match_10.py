import openpyxl

wb = openpyxl.load_workbook(r"D:\KULIAH\Joki\Annisa\skripsi\BAB IV\daftar_pd-SMP TUNAS BANGSA-2026-10-17 05_49_36.xlsx")
ws = wb.active

target_10 = [
    "rahmat padil peratama",
    "abizar Nur ramadhan",
    "yusup Baihaqi",
    "ketut evi yani",
    "mita sari",
    "bima Syarif al bukhory",
    "abi Ibnu Mas'ud",
    "gede radit saputra",
    "reza agus pertama",
    "nyoman mahesa purna"
]

all_rows = []
for r in range(7, ws.max_row + 1):
    nama = ws.cell(r, 2).value
    if nama:
        all_rows.append((r, str(nama).strip()))

print("Matching targets to rows:")
for idx, target in enumerate(target_10, 1):
    t_clean = target.lower().replace("'", "").replace(" ", "").replace("h", "").replace("f", "v").replace("p", "v").replace("z", "s")
    matches = []
    for r, nama in all_rows:
        n_clean = nama.lower().replace("'", "").replace(" ", "").replace("h", "").replace("f", "v").replace("p", "v").replace("z", "s")
        if target.lower() in nama.lower() or nama.lower() in target.lower():
            matches.append((r, nama, 1.0))
        elif t_clean in n_clean or n_clean in t_clean:
            matches.append((r, nama, 0.9))
    if matches:
        matches.sort(key=lambda x: -x[2])
        best = matches[0]
        print(f"{idx}. Target: '{target}' => Row {best[0]}: '{best[1]}'")
    else:
        print(f"{idx}. Target: '{target}' => NOT FOUND")
