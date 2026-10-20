import docx
import sys
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r'D:\KULIAH\Joki\Annisa\skripsi\Annisaaaa.docx'
doc = docx.Document(docx_path)

print(f"Loaded document: {docx_path}")
print(f"Total paragraphs: {len(doc.paragraphs)}")

# Locate P503 (Merode Waterfall)
target_idx = None
for i, p in enumerate(doc.paragraphs):
    if 'Merode Waterfall' in p.text or ('Waterfall' in p.text and p.style.name == 'Heading 2'):
        target_idx = i
        break

print(f"Target heading found at P{target_idx}: {doc.paragraphs[target_idx].text}")

# Fix typo in heading
p_heading = doc.paragraphs[target_idx]
p_heading.text = "Metode Waterfall"
for r in p_heading.runs:
    r.font.name = "Times New Roman"

paras_waterfall = [
    "Menurut Wahid (2020), metode Waterfall merupakan salah satu model klasik dalam System Development Life Cycle (SDLC) yang menerapkan pendekatan sekuensial, linier, dan sistematis dalam rekayasa perangkat lunak. Pada model ini, setiap tahapan pengembangan dirancang mengalir ke bawah seperti air terjun, di mana suatu fase baru dapat dimulai setelah fase sebelumnya diselesaikan secara menyeluruh dan terverifikasi. Karakteristik utama dari metode Waterfall adalah penekanannya pada kejelasan spesifikasi kebutuhan sejak awal proses serta kelengkapan dokumentasi pada setiap tahapan, sehingga meminimalkan ambiguitas dan menjaga konsistensi perancangan sistem informasi berskala terstruktur.",
    
    "Dalam penerapannya, Sasmito (2017) menjelaskan bahwa metode Waterfall terdiri atas lima fase berurutan yang saling berkaitan, yaitu analisis kebutuhan sistem (system requirement analysis), perancangan sistem (design), penulisan kode program (coding/implementation), pengujian program (testing), serta penerapan dan pemeliharaan (implementation and maintenance). Keberhasilan metode ini ditentukan oleh kepatuhan pengembang terhadap urutan proses secara disiplin, di mana hasil luaran (deliverables) dari tahap perancangan logika dan basis data menjadi masukan utama bagi tahap pengkodean program, yang selanjutnya diverifikasi secara menyeluruh menggunakan metode pengujian fungsional sebelum sistem dioperasikan oleh pengguna akhir.",
    
    "Berdasarkan kedua pandangan teoritis tersebut, dapat dipahami bahwa metode Waterfall merupakan metodologi pengembangan perangkat lunak terstruktur yang sangat mengutamakan kepastian alur kerja, kelengkapan analisis data, dan ketepatan tahapan perancangan sebelum memasuki tahap pengkodean. Dalam konteks pembangunan Sistem Pendukung Keputusan Penentuan Prioritas Penerima Program Indonesia Pintar (PIP) berbasis web di SMP Tunas Bangsa, pendekatan Waterfall sangat relevan dan tepat diaplikasikan karena seluruh variabel keputusan—mulai dari lima kriteria penilaian, penetapan skala subkriteria, formulasi matematis algoritma Analytical Hierarchy Process (AHP), alur verifikasi pendaftar, hingga batas kuota penetapan 20 siswa—telah didefinisikan secara pasti dan rigid pada tahap awal penelitian. Dengan mengikuti siklus linier ini, proses translasi dari model konseptual AHP menuju perwujudan sistem web terbebas dari kesalahan logika komputasi, menjamin akurasi perhitungan skor preferensi, serta menghasilkan sistem informasi yang andal, transparan, dan teruji secara fungsional."
]

def format_academic_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Inches(0.4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

# Insert the 3 paragraphs right below the heading
# Next element after heading is doc.paragraphs[target_idx + 1]._p
next_elem = doc.paragraphs[target_idx + 1]._p

for text in paras_waterfall:
    new_p = doc.add_paragraph(text)
    format_academic_paragraph(new_p)
    next_elem.addprevious(new_p._p)

# Remove the empty paragraph that was at target_idx + 1 if empty
if not doc.paragraphs[target_idx + 1 + len(paras_waterfall)].text.strip():
    empty_p = doc.paragraphs[target_idx + 1 + len(paras_waterfall)]._p
    empty_p.getparent().remove(empty_p)

# Save to UPDATED copy
updated_path = r'D:\KULIAH\Joki\Annisa\skripsi\Annisaaaa - UPDATED.docx'
doc.save(updated_path)
print(f"SUCCESS: Saved updated document to: {updated_path}")

# Attempt saving to original
try:
    doc.save(docx_path)
    print(f"SUCCESS: Also overwritten original at: {docx_path}")
except Exception as e:
    print(f"Original file is currently open in Word: {e}")

