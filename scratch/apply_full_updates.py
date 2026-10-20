import docx
import copy
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

docx_path = r'D:\KULIAH\Joki\Annisa\(ok) reNEWok Proposal BAB I-III Annisa belum fix 2 (1) (1).docx'
doc = docx.Document(docx_path)

print(f"Loaded document. Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

# ==============================================================================
# 1. FIX BAB I: Paragraph 415 (Latar Belakang kuota 27 -> 20)
# ==============================================================================
p415 = doc.paragraphs[415]
for r in p415.runs:
    if "27" in r.text:
        print("Fixing P415 run:", repr(r.text))
        r.text = r.text.replace("hanya sebanyak 27", "sebanyak 20").replace("sebanyak 27", "sebanyak 20")

# ==============================================================================
# 2. FIX BAB III: Sync weights percentages in P702-P706
# ==============================================================================
for p in doc.paragraphs[695:715]:
    for r in p.runs:
        if "41.65%" in r.text:
            r.text = r.text.replace("41.65%", "41.62%")
        if "26.19%" in r.text:
            r.text = r.text.replace("26.19%", "26.18%")
        if "16.08%" in r.text:
            r.text = r.text.replace("16.08%", "16.11%")
        if "9.85%" in r.text:
            r.text = r.text.replace("9.85%", "9.86%")
        if "6.23%" in r.text:
            r.text = r.text.replace("6.23%", "6.24%")

# ==============================================================================
# 3. FIX BAB IV: Table Titles & In-text references
# ==============================================================================
# Table 4.12 title (P966)
p966 = doc.paragraphs[966]
for r in p966.runs:
    if r.text == '10':
        r.text = '12'
    if ' 24 ' in r.text:
        r.text = r.text.replace(' 24 ', ' 30 ')

# Table 4.13 title (P973)
p973 = doc.paragraphs[973]
for r in p973.runs:
    if r.text == '11':
        r.text = '13'

# Table 4.14 title (P977)
p977 = doc.paragraphs[977]
for r in p977.runs:
    if r.text == '12':
        r.text = '14'

# Table 4.15 title (P1062)
p1062 = doc.paragraphs[1062]
for r in p1062.runs:
    if r.text == '13':
        r.text = '15'

# In-text references to Table 4.20 in P1061 and P1063
p1061 = doc.paragraphs[1061]
for r in p1061.runs:
    if '4.20' in r.text:
        r.text = r.text.replace('4.20', '4.15')

p1063 = doc.paragraphs[1063]
for r in p1063.runs:
    if '4.20' in r.text:
        r.text = r.text.replace('4.20', '4.15')

# ==============================================================================
# 4. FIX BAB IV: Dashboard Admin (P1027) & Data Calon (P1031) text
# ==============================================================================
p1027 = doc.paragraphs[1027]
for r in p1027.runs:
    if ' (24 ' in r.text:
        r.text = r.text.replace(' (24 ', ' (30 ')
    if ' PIP (10 ' in r.text:
        r.text = r.text.replace(' PIP (10 ', ' PIP (20 ')

p1031 = doc.paragraphs[1031]
for r in p1031.runs:
    if ' (24 ' in r.text:
        r.text = r.text.replace(' (24 ', ' (30 ')

# ==============================================================================
# 5. FIX BAB IV: Figure Captions Shifting (P1010 .. P1058)
# ==============================================================================
fig_map = {
    1010: "Gambar 4. 2 Tampilan Halaman Pendaftaran Mandiri Calon Penerima PIP.",
    1015: "Gambar 4. 3 Tampilan Halaman Login Siswa.",
    1020: "Gambar 4. 4 Tampilan Halaman Pengumuman dan Bukti Pendaftaran Siswa.",
    1025: "Gambar 4. 5 Tampilan Halaman Login Administrator.",
    1029: "Gambar 4. 6 Tampilan Halaman Dashboard Administrator.",
    1033: "Gambar 4. 7 Tampilan Halaman Data Calon Penerima.",
    1038: "Gambar 4. 8 Tampilan Halaman Data Kriteria Penilaian AHP.",
    1043: "Gambar 4. 9 Tampilan Halaman Perhitungan Metode AHP.",
    1048: "Gambar 4. 10 Tampilan Halaman Hasil Ranking Penerima PIP.",
    1053: "Gambar 4. 11 Tampilan Halaman Laporan dan Cetak Dokumen PIP.",
    1058: "Gambar 4. 12 Tampilan Halaman Profil dan Identitas Sekolah."
}

for p_idx, new_caption in fig_map.items():
    p = doc.paragraphs[p_idx]
    p.text = new_caption
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = docx.shared.Pt(11)
        r.font.bold = False

# ==============================================================================
# 6. FIX TABLE 21: Standardize student 18 name
# ==============================================================================
t21 = doc.tables[20] # Tabel 4.12
cell_name_18 = t21.rows[18].cells[2]
if "M. Khian Romadhoni" in cell_name_18.text:
    cell_name_18.text = "Muhammad Khian Romadhoni"
    for r in cell_name_18.paragraphs[0].runs:
        r.font.name = "Times New Roman"
        r.font.size = docx.shared.Pt(10)

# ==============================================================================
# 7. FIX TABLE 22 (Tabel 4.13): Replace 24 old rows with 30 accurate rows
# ==============================================================================
t22 = doc.tables[21]
print(f"Table 22 currently has {len(t22.rows)} rows. Updating to 30 students...")

weights = [0.4162, 0.2618, 0.1611, 0.0986, 0.0624]
students_30 = [
    (1, 'Rahmad Padil Peratama', 5, 5, 5, 5, 5),
    (2, 'Abizar Nur Ramadhan', 5, 5, 5, 4, 5),
    (3, 'Yusup Baihaqi', 5, 5, 4, 5, 5),
    (4, 'Ketut Evi Yani', 5, 5, 4, 4, 5),
    (5, 'Mita Sari', 4, 5, 5, 4, 5),
    (6, 'Bima Syarif Al Bukhory', 4, 5, 4, 4, 5),
    (7, "Abi Ibnu Mas'ud", 4, 4, 4, 3, 4),
    (8, 'Gede Radit Saputra', 4, 4, 3, 3, 4),
    (9, 'Reza Agus Pertama', 3, 4, 4, 3, 4),
    (10, 'Nyoman Mahesa Purna', 3, 4, 3, 3, 4),
    (11, 'Sri Wulandari', 3, 3, 2, 2, 1),
    (12, 'Reza Aditia', 3, 3, 2, 1, 1),
    (13, 'Gea Larassati', 3, 2, 3, 2, 1),
    (14, 'Eka Rohma Wati', 3, 2, 2, 1, 1),
    (15, 'Julia Mulyasari', 3, 2, 1, 1, 2),
    (16, 'Dimas Prayogi', 3, 1, 3, 1, 1),
    (17, 'Wildan Alip Fratama', 1, 3, 3, 2, 2),
    (18, 'Muhammad Khian Romadhoni', 2, 3, 1, 1, 2),
    (19, 'Leka Agustina', 2, 2, 3, 1, 1),
    (20, 'I Komang Mangku Suteja', 3, 1, 1, 2, 2),
    (21, 'Nengah Galih Wijaya', 1, 3, 3, 1, 2),
    (22, 'Rustam', 2, 1, 3, 1, 2),
    (23, 'Nita Sari', 2, 2, 1, 1, 2),
    (24, 'Arif Bastiar', 2, 1, 2, 1, 2),
    (25, 'Mutiara Agung Hasanah', 2, 1, 1, 2, 2),
    (26, 'Alif Fadhil Muzaqi', 1, 3, 1, 1, 1),
    (27, 'Riyan Nofendi', 2, 1, 1, 1, 2),
    (28, 'M Faiq Nur Zaman', 1, 2, 1, 2, 1),
    (29, 'Alfian Nova Pratama', 1, 2, 1, 1, 2),
    (30, 'M. Riski Julian', 1, 1, 1, 2, 1)
]

while len(t22.rows) < 31:
    t22.add_row()

while len(t22.rows) > 31:
    t22._tbl.remove(t22.rows[-1]._tr)

for r_idx, s in enumerate(students_30):
    row = t22.rows[r_idx + 1]
    idx, nama, c1, c2, c3, c4, c5 = s
    v1 = c1 * weights[0]
    v2 = c2 * weights[1]
    v3 = c3 * weights[2]
    v4 = c4 * weights[3]
    v5 = c5 * weights[4]
    tot = v1 + v2 + v3 + v4 + v5
    
    # Cell 0: Clone math A_idx paragraph from Table 21
    src_p = t21.rows[r_idx + 1].cells[0].paragraphs[0]
    c0 = row.cells[0]
    c0._tc.remove(c0.paragraphs[0]._p)
    c0._tc.append(copy.deepcopy(src_p._p))
    
    # Cell 1: Nama
    c1_cell = row.cells[1]
    c1_cell.text = nama
    c1_cell.paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
    for run in c1_cell.paragraphs[0].runs:
        run.font.name = "Times New Roman"
        run.font.size = docx.shared.Pt(10)
        
    # Cells 2-6: Values
    vals = [f"{v1:.4f}".replace('.', ','), f"{v2:.4f}".replace('.', ','), f"{v3:.4f}".replace('.', ','), f"{v4:.4f}".replace('.', ','), f"{v5:.4f}".replace('.', ',')]
    for col_i, val in enumerate(vals):
        cell = row.cells[col_i + 2]
        cell.text = val
        cell.paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = "Times New Roman"
            run.font.size = docx.shared.Pt(10)
            
    # Cell 7: Total Skor
    c7 = row.cells[7]
    c7.text = f"{tot:.4f}".replace('.', ',')
    c7.paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    for run in c7.paragraphs[0].runs:
        run.font.name = "Times New Roman"
        run.font.size = docx.shared.Pt(10)
        run.font.bold = True

print(f"Table 22 successfully updated to {len(t22.rows)} rows.")

# ==============================================================================
# 8. CONDENSE / TIGHTEN BAB IV (4.3.1, 4.3.2, 4.3.3)
# ==============================================================================
# 4.3.1 Software Specification
p983 = doc.paragraphs[983]
p983.text = "Spesifikasi perangkat lunak yang digunakan untuk mendukung proses perancangan, penulisan kode program, pengelolaan basis data, serta pengoperasian Sistem Pendukung Keputusan Penentuan Prioritas Penerima Program Indonesia Pintar (PIP) meliputi sistem operasi Microsoft Windows 11 Home Single Language (64-bit), lingkungan web server lokal Laragon dengan web server Apache 2.4, bahasa pemrograman server-side Hypertext Preprocessor (PHP) versi 8.2 Native guna mengeksekusi logika komputasi AHP, serta MySQL Community Server versi 8.0 sebagai Database Management System (DBMS) pengelola basis data relasional spk_pip. Pada sisi pengembangan kode program, digunakan Visual Studio Code sebagai code editor utama, dengan peramban modern seperti Google Chrome dan Microsoft Edge sebagai media pengujian antarmuka. Sementara itu, perancangan antarmuka pengguna dibangun memanfaatkan kerangka kerja Bootstrap 5 yang didukung oleh pustaka FontAwesome, DataTables untuk pengurutan dan pencarian data dinamis, serta SweetAlert2 untuk penyajian notifikasi dialog interaktif."
p983.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p983.runs:
    r.font.name = "Times New Roman"
    r.font.size = docx.shared.Pt(12)

# 4.3.2 Hardware Specification
p992 = doc.paragraphs[992]
p992.text = "Spesifikasi perangkat keras yang digunakan dalam proses pengembangan, pengujian komputasi algoritma Analytical Hierarchy Process (AHP), dan pengoperasian sistem pendukung keputusan ini didukung oleh spesifikasi komputer berspesifikasi prosesor AMD Athlon Silver 3050U with Radeon Graphics berkecepatan 2.3 GHz (setara minimal Intel Core i3 atau AMD Ryzen 3), didukung oleh Random Access Memory (RAM) berkapasitas 12 GB DDR4 untuk menjamin kelancaran multitasking pemrosesan komputasi data, serta media penyimpanan utama berupa Solid State Drive (SSD) berkapasitas 256 GB guna mempercepat proses baca dan tulis data pada sistem basis data. Untuk visualisasi dan interaksi antarmuka pengguna, sistem dioperasikan menggunakan monitor berukuran 14 inci beresolusi Full HD (1920 \u00d7 1080 piksel), serta dilengkapi perangkat masukan dan keluaran berupa keyboard, mouse/touchpad, dan printer untuk pencetakan dokumen fisik rekomendasi penetapan kuota PIP."
p992.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p992.runs:
    r.font.name = "Times New Roman"
    r.font.size = docx.shared.Pt(12)

# 4.3.3 Database Specification
p1001 = doc.paragraphs[1001]
p1001.text = "Kelima tabel utama tersebut memiliki fungsi fungsional yang saling terintegrasi dalam mendukung operasional sistem pendukung keputusan. Tabel admin berfungsi menyimpan data akun dan kredensial login administrator sistem yang terenkripsi. Tabel pengaturan digunakan untuk menyimpan parameter konfigurasi identitas kelembagaan sekolah, alokasi batas kuota penerima PIP (20 siswa), periode pembukaan dan penutupan pendaftaran mandiri, serta nilai parameter konsistensi AHP. Tabel calon_penerima berfungsi sebagai entitas data master seluruh pendaftar, memuat identitas peserta didik, pemetaan nilai skala subkriteria (C1 hingga C5), hasil akumulasi skor akhir AHP, urutan peringkat prioritas, dan status verifikasi berkas. Sementara itu, tabel kriteria dan tabel matriks_kriteria bertindak sebagai media penyimpan data konfigurasi metode AHP, mencakup definisi lima kriteria penilaian, sifat kriteria (cost/benefit), bobot prioritas vektor eigen hasil kalkulasi, serta nilai intensitas perbandingan berpasangan antar kriteria pada skala Saaty 1\u20139."
p1001.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p1001.runs:
    r.font.name = "Times New Roman"
    r.font.size = docx.shared.Pt(12)

# Safely remove redundant list paragraphs in reverse order
paras_to_remove = [
    # 4.3.3 database list items: 1006, 1005, 1004, 1003, 1002
    1006, 1005, 1004, 1003, 1002,
    # 4.3.2 hardware list items: 997, 996, 995, 994, 993
    997, 996, 995, 994, 993,
    # 4.3.1 software list items: 990, 989, 988, 987, 986, 985, 984
    990, 989, 988, 987, 986, 985, 984
]

for idx in sorted(paras_to_remove, reverse=True):
    p_elem = doc.paragraphs[idx]._p
    p_elem.getparent().remove(p_elem)
    print(f"Removed redundant list paragraph P{idx}")

# Save updated document
output_path_updated = r'D:\KULIAH\Joki\Annisa\(ok) reNEWok Proposal BAB I-III Annisa belum fix 2 (1) (1) - UPDATED.docx'
doc.save(output_path_updated)
print(f"SUCCESS: Saved updated document to: {output_path_updated}")

try:
    doc.save(docx_path)
    print(f"SUCCESS: Also overwritten original document at: {docx_path}")
except Exception as e:
    print(f"NOTE: Original document is currently open in Microsoft Word. It has been saved to: {output_path_updated}")

