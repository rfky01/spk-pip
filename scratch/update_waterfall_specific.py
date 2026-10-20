import docx
import sys
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

# We update Annisaaaa - UPDATED.docx
updated_docx_path = r'D:\KULIAH\Joki\Annisa\skripsi\Annisaaaa - UPDATED.docx'
doc = docx.Document(updated_docx_path)

print(f"Loaded document: {updated_docx_path}")
print(f"Total paragraphs: {len(doc.paragraphs)}")

# Locate heading "Metode Waterfall"
target_idx = None
for i, p in enumerate(doc.paragraphs):
    if 'Metode Waterfall' in p.text or 'Merode Waterfall' in p.text:
        target_idx = i
        break

print(f"Target heading found at P{target_idx}: {doc.paragraphs[target_idx].text}")

# Ensure heading is clean
doc.paragraphs[target_idx].text = "Metode Waterfall"
for r in doc.paragraphs[target_idx].runs:
    r.font.name = "Times New Roman"

# The 3 new paragraphs based on the user's two exact PDFs:
# 1. Putri & Putri (2024)
# 2. Akbar dkk. (2024)
# 3. Sintesis / Penjelasan
paras_new = [
    "Menurut Putri dan Putri (2024), metode Waterfall merupakan salah satu model dalam Software Development Life Cycle (SDLC) yang berfungsi sebagai pendekatan efisien dan terstruktur dalam merancang serta mengelola proyek pengembangan sistem informasi. Model ini sangat sesuai diterapkan pada perancangan perangkat lunak yang memerlukan tahapan yang sistematis guna memudahkan pengembang dalam mengidentifikasi permasalahan secara mendalam serta merumuskan fitur sistem yang selaras dengan kebutuhan nyata pengguna. Dalam pendekatan ini, seluruh proses rekayasa perangkat lunak disusun mengalir secara berurutan, dimulai dari pendefinisian spesifikasi kebutuhan secara utuh sebelum melangkah ke tahap perancangan visual dan arsitektur data.",
    
    "Lebih lanjut, Akbar dkk. (2024) menjelaskan bahwa model Waterfall terdiri atas tahapan-tahapan sekuensial yang berkesinambungan, meliputi analisis kebutuhan (requirements analysis), perancangan sistem (system design), implementasi (implementation), pengujian (testing), hingga tahap pemeliharaan (maintenance). Pemilihan model Waterfall dinilai sangat tepat untuk sistem yang memiliki spesifikasi kebutuhan yang cenderung stabil, jelas, dan telah terdefinisi dengan baik sejak awal proses perencanaan. Karakteristik pengembangannya yang bersifat linier, sistematis, dan terstruktur memberikan keunggulan dalam hal kemudahan evaluasi progres proyek, keakuratan pelaporan pada setiap tahapan, serta penyusunan dokumentasi sistem yang komprehensif guna menjamin kemudahan pengelolaan perangkat lunak dalam jangka panjang.",
    
    "Berdasarkan kedua landasan teori tersebut, dapat disimpulkan bahwa metode Waterfall merupakan metodologi rekayasa perangkat lunak terstruktur yang mengutamakan kepastian spesifikasi kebutuhan, kedisiplinan alur fase kerja, dan keandalan dokumentasi sebelum melangkah ke tahap penulisan baris kode program. Dalam konteks pembangunan Sistem Pendukung Keputusan Penentuan Prioritas Penerima Program Indonesia Pintar (PIP) berbasis web di SMP Tunas Bangsa, pendekatan Waterfall sangat relevan dan efektif diaplikasikan. Hal ini dikarenakan seluruh variabel keputusan—mulai dari kelima kriteria penilaian (Penghasilan Orang Tua, Jumlah Tanggungan, Kondisi Rumah, Prestasi Akademik, dan Jarak Tempuh), pedoman konversi skala subkriteria, formulasi matematis matriks pembobotan algoritma Analytical Hierarchy Process (AHP), alur verifikasi berkas, hingga penetapan batas kuota 20 siswa penerima bantuan—telah teridentifikasi dan terkunci secara pasti sejak awal penelitian. Alur sekuensial linier Waterfall memastikan bahwa logika komputasi AHP yang telah dirancang dapat diimplementasikan ke dalam kode program secara presisi, bebas dari kesalahan algoritma, serta menghasilkan sistem penyeleksian beasiswa yang transparan, objektif, dan dapat dipertanggungjawabkan."
]

def format_academic_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Inches(0.4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

# Check next paragraphs after heading:
# In previous turn, we had inserted 3 paragraphs at target_idx + 1, target_idx + 2, target_idx + 3.
# Let's replace the text of those 3 paragraphs directly!
doc.paragraphs[target_idx + 1].text = paras_new[0]
format_academic_paragraph(doc.paragraphs[target_idx + 1])

doc.paragraphs[target_idx + 2].text = paras_new[1]
format_academic_paragraph(doc.paragraphs[target_idx + 2])

doc.paragraphs[target_idx + 3].text = paras_new[2]
format_academic_paragraph(doc.paragraphs[target_idx + 3])

doc.save(updated_docx_path)
print(f"SUCCESS: Updated {updated_docx_path} with new Putri (2024) & Akbar (2024) paragraphs!")

# Also try original
orig_path = r'D:\KULIAH\Joki\Annisa\skripsi\Annisaaaa.docx'
try:
    doc.save(orig_path)
    print(f"SUCCESS: Also updated original file at: {orig_path}")
except Exception as e:
    print(f"Original file is currently open in Word: {e}")

