import openpyxl

wb = openpyxl.load_workbook(r"D:\KULIAH\Joki\Annisa\skripsi\BAB IV\Data_Calon_Penerima_PIP_SMP_TUNAS_BANGSA.xlsx")
ws = wb.active

print("Headers (Row 4):")
headers = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
print(headers)

print("\nTop 10 Data Rows:")
for r in range(5, 15):
    row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    no = row_vals[0]
    nisn = row_vals[1]
    nama = row_vals[2]
    ortu = row_vals[3]
    hp = row_vals[4]
    c1 = row_vals[9]
    print(f"{no}. {nama} | Ortu: {ortu} | NISN: {nisn} | HP: {hp} | Penghasilan: {c1}")
