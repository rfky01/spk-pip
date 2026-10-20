import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r'D:\KULIAH\Joki\Annisa\(ok) Skripsi BAB I-V Annisa (LENGKAP).docx'
doc = docx.Document(docx_path)

print(f"Loaded: {docx_path}")
print(f"Total Paragraphs: {len(doc.paragraphs)}, Total Tables: {len(doc.tables)}")

# Print all paragraphs in BAB V
print("\n=== BAB V COMPLETE INSPECTION ===")
for i, p in enumerate(doc.paragraphs):
    if i >= 1025:
        print(f"P{i} [{p.style.name}]: {p.text}")

# Check Table 22 in final file
t22 = doc.tables[21]
print(f"\nTable 22 rows: {len(t22.rows)}")
print(f"Table 22 Row 1: {[c.text.strip() for c in t22.rows[1].cells]}")
print(f"Table 22 Row 20: {[c.text.strip() for c in t22.rows[20].cells]}")
print(f"Table 22 Row 30: {[c.text.strip() for c in t22.rows[30].cells]}")

