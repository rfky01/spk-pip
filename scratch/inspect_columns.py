import openpyxl

wb = openpyxl.load_workbook(r"D:\KULIAH\Joki\Annisa\skripsi\BAB IV\daftar_pd-SMP TUNAS BANGSA-2026-10-17 05_49_36.xlsx")
ws = wb.active

print("Headers in row 5:")
for c in [5, 2, 20, 4, 57, 10, 11, 12, 13, 14, 15, 29, 35, 17, 66]:
    print(f"Col {c}: {ws.cell(5, c).value} (r6: {ws.cell(6, c).value})")

# Look at some sample rows
print("\nSample rows:")
for r in [7, 8, 9, 25, 46, 56, 71, 84, 91, 98, 113]:
    nama = ws.cell(r, 2).value
    nisn = ws.cell(r, 5).value
    hp = ws.cell(r, 20).value
    jk = ws.cell(r, 4).value
    asal = ws.cell(r, 57).value
    alamat = ws.cell(r, 10).value
    dusun = ws.cell(r, 13).value
    desa = ws.cell(r, 14).value
    kec = ws.cell(r, 15).value
    gaji_ayah = ws.cell(r, 29).value
    gaji_ibu = ws.cell(r, 35).value
    tinggal = ws.cell(r, 17).value
    jarak = ws.cell(r, 66).value
    print(f"Row {r}: {nama} | NISN:{nisn} | HP:{hp} | JK:{jk} | Asal:{asal} | Alamat:{alamat}, {dusun}, {desa} | GajiA:{gaji_ayah} | Tinggal:{tinggal} | Jarak:{jarak}")
