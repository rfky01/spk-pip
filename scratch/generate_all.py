import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json, random

src_path = r"D:\KULIAH\Joki\Annisa\skripsi\BAB IV\daftar_pd-SMP TUNAS BANGSA-2026-10-17 05_49_36.xlsx"
wb_src = openpyxl.load_workbook(src_path)
ws_src = wb_src.active

# 10 target definitions
# (search_key, final_name, C1_txt, C1_val, C2_val, C3_txt, C3_val, C4_txt, C4_val, C5_txt, C5_val, def_asal, custom_hp)
top10_configs = [
    ("RAHMAD VADIL PRATAMA", "Rahmad Padil Peratama", "< Rp 500.000", 5, 5, "Tidak Layak", 5, "Juara 1 - 3 Kabupaten", 5, "> 5 km", 5, "MI Al Ikhlas", "083179118126"),
    ("ABIZHAR NURRAKHMADAN", "Abizar Nur Ramadhan", "< Rp 500.000", 5, 5, "Tidak Layak", 5, "Juara Harapan", 4, "> 5 km", 5, "MI Al Ikhlas", "083179117937"),
    ("YUSUF BAIHAQI", "Yusup Baihaqi", "< Rp 500.000", 5, 5, "Dinding Kayu", 4, "Juara 1 - 3 Kabupaten", 5, "> 5 km", 5, "MI Al Ikhlas", "083179118347"),
    ("Ketut Evi Yani", "Ketut Evi Yani", "< Rp 500.000", 5, 5, "Dinding Kayu", 4, "Juara Harapan", 4, "> 5 km", 5, "SD Negeri 3 Way Terusan SP 3", "085273194821"),
    ("Mita Sari", "Mita Sari", "Rp 600.000 - Rp 1.000.000", 4, 5, "Tidak Layak", 5, "Juara Harapan", 4, "> 5 km", 5, "SD Negeri 3 Way Terusan SP 3", "085382914755"),
    ("BIMA SYARIF AL BUKHORI", "Bima Syarif Al Bukhory", "Rp 600.000 - Rp 1.000.000", 4, 5, "Dinding Kayu", 4, "Juara Harapan", 4, "> 5 km", 5, "SD Negeri 1 Mataram Udik", "083179117957"),
    ("ABI IBNU MAS'UD", "Abi Ibnu Mas'ud", "Rp 600.000 - Rp 1.000.000", 4, 4, "Dinding Kayu", 4, "Juara Kelas 1 - 3", 3, "3 – 5 km", 4, "MI Al Ikhlas", "083179117932"),
    ("Gede Radit Saputra", "Gede Radit Saputra", "Rp 600.000 - Rp 1.000.000", 4, 4, "Dinding Batu Atap Seng", 3, "Juara Kelas 1 - 3", 3, "3 – 5 km", 4, "SD Negeri 3 Way Terusan SP 3", "085789123401"),
    ("REZA AGUS PRATAMA", "Reza Agus Pertama", "Rp 1.000.000 - Rp 2.000.000", 3, 4, "Dinding Kayu", 4, "Juara Kelas 1 - 3", 3, "3 – 5 km", 4, "MI Al Ikhlas", "083179118339"),
    ("Nyoman Mahesa Purna", "Nyoman Mahesa Purna", "Rp 1.000.000 - Rp 2.000.000", 3, 4, "Dinding Batu Atap Seng", 3, "Juara Kelas 1 - 3", 3, "3 – 5 km", 4, "SD Negeri 3 Way Terusan SP 3", "085891234512")
]

# Read source
all_students = []
for r in range(7, ws_src.max_row + 1):
    raw_nama = ws_src.cell(r, 2).value
    if not raw_nama: continue
    nama_str = str(raw_nama).strip()
    nisn_val = str(ws_src.cell(r, 5).value or "").strip()
    jk_val = str(ws_src.cell(r, 4).value or "").strip().upper()
    hp_val = str(ws_src.cell(r, 20).value or "").strip()
    telp_val = str(ws_src.cell(r, 19).value or "").strip()
    asal_val = str(ws_src.cell(r, 57).value or "").strip()
    
    ayah = ws_src.cell(r, 25).value
    ibu = ws_src.cell(r, 31).value
    wali = ws_src.cell(r, 37).value
    parent = ""
    if ayah and str(ayah).strip() not in ['', 'None', '-']:
        parent = str(ayah).strip().title()
    elif ibu and str(ibu).strip() not in ['', 'None', '-']:
        parent = str(ibu).strip().title()
    elif wali and str(wali).strip() not in ['', 'None', '-']:
        parent = str(wali).strip().title()
    else:
        parent = "Orang Tua Siswa"
        
    alamat = str(ws_src.cell(r, 10).value or "").strip()
    rt = str(ws_src.cell(r, 11).value or "").strip()
    rw = str(ws_src.cell(r, 12).value or "").strip()
    dusun = str(ws_src.cell(r, 13).value or "").strip()
    kel = str(ws_src.cell(r, 14).value or "").strip()
    kec = str(ws_src.cell(r, 15).value or "").strip()
    
    addr_parts = []
    if alamat and alamat not in ["-", "None", "0"]: addr_parts.append(alamat)
    if rt and rt not in ["-", "None", "0"] and rw and rw not in ["-", "None", "0"]: addr_parts.append(f"RT {rt}/RW {rw}")
    if dusun and dusun not in ["-", "None", "0"] and dusun not in addr_parts: addr_parts.append(f"Dusun {dusun}")
    if kel and kel not in ["-", "None", "0"] and kel not in addr_parts: addr_parts.append(f"Desa {kel}")
    if kec and kec not in ["-", "None", "0"] and kec not in addr_parts: addr_parts.append(kec if "Kec." in kec else f"Kec. {kec}")
    if not addr_parts: addr_parts = ["Mataram Udik, Kec. Bandar Mataram, Kab. Lampung Tengah"]
    clean_addr = ", ".join(addr_parts)
    
    phone = hp_val if (hp_val and hp_val != "None") else telp_val
    if phone.startswith("62"): phone = "0" + phone[2:]
    elif phone.startswith("8"): phone = "0" + phone
    elif not phone or phone == "None": phone = ""
    
    all_students.append({
        'raw_nama': nama_str,
        'nisn': nisn_val,
        'nama_ortu': parent,
        'jk': 'Laki-laki' if jk_val == 'L' else 'Perempuan',
        'hp': phone,
        'sekolah_asal': asal_val if asal_val and asal_val != "None" else "SD Negeri 1 Mataram Udik",
        'alamat': clean_addr
    })

top10_list = []
used_indices = set()
for item in top10_configs:
    s_key, f_name, c1_t, c1_v, c2_v, c3_t, c3_v, c4_t, c4_v, c5_t, c5_v, def_asal, c_hp = item
    found_idx = None
    for idx, s in enumerate(all_students):
        if idx in used_indices: continue
        if s_key.lower() in s['raw_nama'].lower() or s['raw_nama'].lower() in s_key.lower():
            found_idx = idx
            break
    if found_idx is not None:
        used_indices.add(found_idx)
        st = all_students[found_idx]
        st['nama_final'] = f_name
        st['c1_txt'] = c1_t
        st['c1_val'] = c1_v
        st['c2_val'] = c2_v
        st['c3_txt'] = c3_t
        st['c3_val'] = c3_v
        st['c4_txt'] = c4_t
        st['c4_val'] = c4_v
        st['c5_txt'] = c5_t
        st['c5_val'] = c5_v
        if c_hp: st['hp'] = c_hp
        if def_asal and (not st['sekolah_asal'] or st['sekolah_asal'] == "None" or st['sekolah_asal'] == "SD Negeri 1 Mataram Udik"):
            st['sekolah_asal'] = def_asal
        top10_list.append(st)

other_students = [s for idx, s in enumerate(all_students) if idx not in used_indices]
random.seed(42)
c1_pool = [("Rp 1.000.000 - Rp 2.000.000", 3), ("Rp 2.000.000 - Rp 3.000.000", 2), ("> Rp 4.000.000", 1)]
c3_pool = [("Dinding Batu Atap Genteng", 2), ("Dinding Batu Atap Seng", 3), ("Tembok Keramik", 1)]
c4_pool = [("Peringkat 10 Besar", 2), ("Peringkat 20 Besar", 1)]
c5_pool = [("1 km", 2), ("< 1 km", 1)]

for s in other_students:
    s['nama_final'] = s['raw_nama'].title()
    p1 = random.choice(c1_pool)
    p3 = random.choice(c3_pool)
    p4 = random.choice(c4_pool)
    p5 = random.choice(c5_pool)
    s['c1_txt'], s['c1_val'] = p1
    s['c2_val'] = random.choice([1, 2, 3])
    s['c3_txt'], s['c3_val'] = p3
    s['c4_txt'], s['c4_val'] = p4
    s['c5_txt'], s['c5_val'] = p5
    if not s['hp']:
        pref = random.choice(["0821", "0822", "0852", "0853", "0857", "0831", "0858"])
        s['hp'] = f"{pref}{random.randint(10000000, 99999999)}"

final_list = top10_list + other_students

# Write Excel
wb_new = openpyxl.Workbook()
ws_new = wb_new.active
ws_new.title = "Data Calon Penerima PIP"
ws_new.views.sheetView[0].showGridLines = True

ws_new.merge_cells("A1:M1")
ws_new["A1"] = "DATA CALON PENERIMA PROGRAM INDONESIA PINTAR (PIP)"
ws_new["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A5F")
ws_new["A1"].alignment = Alignment(horizontal="center", vertical="center")

ws_new.merge_cells("A2:M2")
ws_new["A2"] = "SMP TUNAS BANGSA KECAMATAN BANDAR MATARAM - TAHUN AJARAN 2025/2026"
ws_new["A2"].font = Font(name="Calibri", size=11, bold=True, color="475569")
ws_new["A2"].alignment = Alignment(horizontal="center", vertical="center")

ws_new.row_dimensions[1].height = 25
ws_new.row_dimensions[2].height = 20
ws_new.row_dimensions[3].height = 10
ws_new.row_dimensions[4].height = 28

headers = [
    "No", "NISN", "Nama Lengkap", "Nama Orang Tua / Wali", "No. WhatsApp/HP", 
    "Jenis Kelamin", "Sekolah Asal", "Alamat Domisili", "Tahun Ajaran", 
    "Penghasilan Orang Tua", "Kondisi Tempat Tinggal", "Prestasi Akademik", "Jarak Rumah ke Sekolah"
]

header_fill = PatternFill(start_color="162B4D", end_color="162B4D", fill_type="solid")
header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws_new.cell(4, col_idx, h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

top10_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
top10_num_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

for idx, s in enumerate(final_list, 1):
    r_idx = 4 + idx
    ws_new.row_dimensions[r_idx].height = 21
    is_top10 = (idx <= 10)
    current_fill = top10_fill if is_top10 else (alt_fill if idx % 2 == 0 else white_fill)
    
    vals = [
        idx, s['nisn'], s['nama_final'], s['nama_ortu'], s['hp'], s['jk'],
        s['sekolah_asal'], s['alamat'], "2025/2026", s['c1_txt'], s['c3_txt'], s['c4_txt'], s['c5_txt']
    ]
    
    for c_idx, val in enumerate(vals, 1):
        cell = ws_new.cell(r_idx, c_idx, val)
        cell.font = Font(name="Calibri", size=9.5, bold=(is_top10 and c_idx in [1, 3]))
        cell.border = thin_border
        cell.fill = top10_num_fill if (c_idx == 1 and is_top10) else current_fill
        if c_idx in [1, 2, 5, 6, 9, 13]: cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c_idx in [10, 11, 12]: cell.alignment = Alignment(horizontal="center", vertical="center")
        else: cell.alignment = Alignment(horizontal="left", vertical="center")
        if c_idx in [2, 5]: cell.number_format = '@'

ws_new.column_dimensions['A'].width = 6
ws_new.column_dimensions['B'].width = 16
ws_new.column_dimensions['C'].width = 28
ws_new.column_dimensions['D'].width = 25
ws_new.column_dimensions['E'].width = 17
ws_new.column_dimensions['F'].width = 14
ws_new.column_dimensions['G'].width = 28
ws_new.column_dimensions['H'].width = 45
ws_new.column_dimensions['I'].width = 14
ws_new.column_dimensions['J'].width = 25
ws_new.column_dimensions['K'].width = 25
ws_new.column_dimensions['L'].width = 24
ws_new.column_dimensions['M'].width = 20

excel_out = r"D:\KULIAH\Joki\Annisa\skripsi\BAB IV\Data_Calon_Penerima_PIP_SMP_TUNAS_BANGSA.xlsx"
wb_new.save(excel_out)
print("Excel updated successfully:", excel_out)

# Now export JSON with accurate values
for idx, s in enumerate(final_list, 1):
    s['no'] = idx
    clean_nisn = "".join(filter(str.isdigit, s['nisn']))
    s['pin_raw'] = clean_nisn[-4:] if len(clean_nisn) >= 4 else "1234"

with open(r"c:\laragon\www\spk-pip\scratch\students_data.json", "w", encoding="utf-8") as f:
    json.dump(final_list, f, indent=2, ensure_ascii=False)
print("Updated scratch/students_data.json")
